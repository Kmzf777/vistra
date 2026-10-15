"""Gera a planilha de leads classificados (QUENTE / MORNO / DESCARTE).

Os dados ficam fora do git, num JSON em leads/ (ex.: leads/dentista-uberlandia-2026-10-10.json).
Rode a partir de sites/:

    python scripts/gerar_planilha.py leads/<arquivo>.json [saida.xlsx]

Sem a saída, grava o .xlsx ao lado do JSON, com o mesmo nome.

Formato do JSON:
    {
      "regiao": "Uberlândia - MG",
      "fonte": "de onde vieram os dados (scraper, data, busca)",
      "leads": [
        {"nome", "endereco", "telefone", "nota", "avaliacoes", "categoria", "site_maps", "place_id",
         "instagram", "seguidores", "posts", "ultimo_post", "link_bio", "tipo_link_bio",
         "classificacao" (QUENTE | MORNO | DESCARTE), "motivo"}
      ]
    }
A ordem dos MORNOS no JSON é a ordem de prioridade.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

M = "https://www.google.com/maps/search/?api=1&query=Google&query_place_id="

if len(sys.argv) < 2:
    sys.exit(__doc__)
entrada = Path(sys.argv[1])
saida = Path(sys.argv[2]) if len(sys.argv) > 2 else entrada.with_suffix(".xlsx")
dados = json.loads(entrada.read_text(encoding="utf-8"))
L = dados["leads"]

CAB = ["Nome", f"Endereço ({dados.get('regiao', '')})", "Telefone", "Nota", "Avaliações", "Categoria", "Site no Maps",
       "Instagram", "Seguidores", "Posts", "Último post", "Link da bio", "Tipo do link da bio",
       "Motivo da classificação", "Link do Maps"]
LARG = [32, 44, 18, 7, 11, 20, 36, 26, 11, 8, 12, 36, 18, 70, 40]
FILL = {"QUENTE": "F4D7C8", "MORNO": "F3EBD3", "DESCARTE": "E5E5E5"}
LINK = Font(name="Arial", size=10, color="7E5A3E", underline="single")

wb = Workbook()
wb.remove(wb.active)
for classe in ["QUENTE", "MORNO", "DESCARTE"]:
    ws = wb.create_sheet(classe)
    ws.append(CAB)
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E2C2F")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    # QUENTE e DESCARTE por nº de avaliações; MORNO mantém a ordem de prioridade do JSON.
    linhas = [l for l in L if l["classificacao"] == classe]
    if classe != "MORNO":
        linhas.sort(key=lambda l: -l["avaliacoes"])
    for l in linhas:
        ws.append([l["nome"], l["endereco"], l["telefone"], l["nota"], l["avaliacoes"], l["categoria"],
                   l["site_maps"], l["instagram"], l["seguidores"], l["posts"], l["ultimo_post"],
                   l["link_bio"], l["tipo_link_bio"], l["motivo"], M + l["place_id"]])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
        row[3].number_format = "0.0"
        row[0].fill = PatternFill("solid", fgColor=FILL[classe])
        row[0].font = Font(name="Arial", size=10, bold=True)
        for idx in (6, 11, 14):
            if row[idx].value:
                row[idx].hyperlink = row[idx].value
                row[idx].font = LINK
    for i, w in enumerate(LARG, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

ws = wb.create_sheet("Critérios")
for r in [
    ["Critério", "Regra"],
    ["QUENTE", "Sem site próprio (ou o site do Maps é Instagram, WhatsApp ou agregador), 20+ avaliações, nota 4,3+ e post no Instagram nos últimos 30 dias."],
    ["MORNO", "Sem site próprio, mas com poucas avaliações, nota abaixo de 4,3, Instagram parado/privado ou Instagram ainda não verificado."],
    ["DESCARTE", "Já tem site próprio decente, ou parece fechado/abandonado."],
    ["Fonte", dados.get("fonte", "")],
    ["Último post", "Data do post mais recente da própria conta (posts fixados e de outras contas ignorados)."],
]:
    ws.append(r)
for c in ws[1]:
    c.font = Font(name="Arial", bold=True)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name="Arial", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 110

wb.save(saida)
print({k: sum(1 for l in L if l["classificacao"] == k) for k in FILL}, len(L), "->", saida)
