# -*- coding: utf-8 -*-
"""Monta a fila de cold call (site + CRM com API oficial do WhatsApp) em Uberlândia.

Entrada : leads/uberlandia/raw/<nicho>.tsv  (raspagem do Google Maps, uma linha por lugar)
          colunas: nome, nota, avaliacoes, categoria, endereco, telefone, site, fechado, patrocinado
Saída   : leads/uberlandia/leads.json   -> lido pela /lista do Vistra CRM e pelo crm/scripts/importar-leads.js
          leads/uberlandia/coldcall_uberlandia.xlsx + .csv -> planilha para ligar

    python scripts/montar_leads_udi.py            (enriquece visitando os sites)
    python scripts/montar_leads_udi.py --rapido   (sem visitar sites)

O enriquecimento visita só o site que o próprio Google Maps lista para a empresa,
procurando Instagram, e-mail, link de WhatsApp e sinais de chatbot/CRM já instalado.
"""
import csv, hashlib, json, re, sys, unicodedata, urllib.request
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
PASTA = RAIZ / "leads" / "uberlandia"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36"

# arquivo -> (nome do nicho, por que site + CRM paga a conta nesse nicho)
NICHOS = {
    "odonto":      ("Odontologia",        "implante e ortodontia são ticket alto; lead que espera no WhatsApp marca no concorrente"),
    "estetica":    ("Estética",           "harmonização e procedimentos vivem de Instagram + WhatsApp; follow-up é onde o dinheiro escapa"),
    "imobiliaria": ("Imobiliária",        "cada imóvel gera dezenas de conversas; sem CRM o corretor perde o lead quente"),
    "autoescola":  ("Autoescola",         "muita pergunta repetida de preço e pacote; bot oficial responde e agenda"),
    "academia":    ("Academia/Studio",    "aula experimental, plano e renovação: funil perfeito para WhatsApp + CRM"),
    "seminovos":   ("Seminovos",          "lead de carro esfria em horas; quem responde primeiro vende"),
    "planejados":  ("Móveis planejados",  "orçamento longo com várias idas e vindas; CRM segura o funil"),
    "solar":       ("Energia solar",      "venda consultiva com simulação; site captura conta de luz e CRM acompanha"),
    "idiomas":     ("Escola/Cursos",      "matrícula sazonal e muita dúvida; disparo oficial reativa ex-alunos"),
    "veterinaria": ("Veterinária",        "agenda, vacina e retorno: lembrete automático pelo WhatsApp oficial"),
    "petshop":     ("Pet shop",           "banho e tosa recorrente e pedido de ração no WhatsApp; agenda e lembrete automáticos"),
    "advocacia":   ("Advocacia",          "triagem de caso pelo WhatsApp e site com autoridade local"),
    "eventos":     ("Buffet/Eventos",     "orçamento por data, alto ticket, muita conversa até fechar"),
}

# Redes e franquias: a decisão de site/CRM não é da unidade de Uberlândia.
REDES = ["sorridents", "odonto company", "oral sin", "odontoclinic", "uniodonto", "oralmedic",
         "dentistas do trabalhador", "smart fit", "bluefit", "selfit", "pratique", "wizard", "ccaa",
         "fisk", "cna ", "kumon", "open english", "localiza", "movida", "unidas", "todeschini",
         "dell anno", "florense", "italinea", "kitchens", "sim clinica", "rede sim", "petz", "cobasi",
         "espacolaser", "espaço laser", "depyl", "buddemeyer", "o boticario", "jequiti"]

# Público, sistema S e categoria que não compra site + CRM de agência local.
FORA = ["instituto federal", "universidade federal", "ufu ", "senai", "senac", "sesc", "sesi", "sebrae",
        "prefeitura", "faculdade", "detran", "escola estadual", "escola municipal"]
CAT_FORA = ["psicólogo", "peças para motocicletas", "oficina mecânica", "artigos para festas", "coworking",
            "loja de material de construção", "escola pública", "universidade", "órgão"]

SITE_FRACO = ["linktr.ee", "canva.site", "my.canva", "agenda.link", "agendarconsulta", "bio.link",
              "beacons.ai", "wixsite.com", "sites.google", "negocio.site", "business.site",
              "linkbio", "taplink", "lnk.bio", "carrd.co"]

# Sinais de que já existe atendimento automatizado / CRM no site.
CRM = {"kommo": "Kommo", "amocrm": "Kommo", "rdstation": "RD Station", "rd-station": "RD Station",
       "blip.ai": "Blip", "take.net": "Blip", "zenvia": "Zenvia", "octadesk": "Octadesk",
       "huggy": "Huggy", "jivosite": "JivoChat", "jivochat": "JivoChat", "tawk.to": "Tawk",
       "chatwoot": "Chatwoot", "manychat": "ManyChat", "botconversa": "BotConversa",
       "leadster": "Leadster", "hubspot": "HubSpot", "pipedrive": "Pipedrive", "digisac": "Digisac",
       "z-api": "Z-API", "chatguru": "ChatGuru", "poli.digital": "Poli",
       "helena.app": "Helena", "utalk.umbler": "Umbler Talk", "clinicorp": "Clinicorp"}
# Fora da lista de propósito: "sults" casa com "results" e "umbler" é hospedagem.

EMAIL_RUIDO = re.compile(r"(sentry|wixpress|example|seudominio|email@|nome@|@sentry|\.png|\.jpg|\.webp|godaddy|domain\.com)", re.I)


def norm(s):
    s = unicodedata.normalize("NFD", (s or "").lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def baixar(url):
    try:
        r = urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Language": "pt-BR"})
        with urllib.request.urlopen(r, timeout=10) as resp:
            return resp.read(600_000).decode("utf-8", "ignore")
    except Exception:
        return ""


def enriquecer(lead):
    """Visita o site listado no Maps e extrai IG, e-mail, WhatsApp e CRM."""
    html = baixar(lead["site"])
    lead["site_respondeu"] = bool(html)
    if not html:
        return lead
    low = html.lower()
    if not lead["ig"]:
        for m in re.finditer(r"instagram\.com/([A-Za-z0-9_.]{2,30})", html):
            u = m.group(1).rstrip(".")
            if u.lower() not in {"p", "reel", "explore", "accounts", "stories", "share", "tv", "_u"}:
                lead["ig"] = f"https://www.instagram.com/{u}/"
                break
    if not lead["email"]:
        for e in re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", html):
            if not EMAIL_RUIDO.search(e):
                lead["email"] = e.lower()
                break
    lead["wa_no_site"] = bool(re.search(r"wa\.me/|api\.whatsapp\.com|web\.whatsapp\.com", low))
    lead["crm"] = sorted({nome for k, nome in CRM.items() if k in low})
    return lead


def ler_brutos():
    vistos, out = {}, []
    for arq in sorted((PASTA / "raw").glob("*.tsv")):
        chave = arq.stem
        nicho, porque = NICHOS.get(chave, (chave.title(), ""))
        for linha in arq.read_text(encoding="utf-8").splitlines():
            c = (linha.split("\t") + [""] * 9)[:9]
            nome, nota, aval, cat, end, tel, site, fechado, _pat = [x.strip() for x in c]
            if not nome or fechado == "1":
                continue
            if any(r in norm(nome) for r in REDES) or any(f in norm(nome) for f in FORA):
                continue
            if any(norm(f) in norm(cat) for f in CAT_FORA):
                continue
            dig = re.sub(r"\D", "", tel)
            k = (norm(nome)[:28], dig)
            if k in vistos or (dig and dig in vistos):
                continue
            vistos[k] = 1
            if dig:
                vistos[dig] = 1
            ig = fb = ""
            sl = site.lower()
            if "instagram.com" in sl:
                ig, site = site, ""
            elif "facebook.com" in sl:
                fb, site = site, ""
            elif "wa.me" in sl or "whatsapp" in sl:
                site = ""
            out.append({
                "id": "udi-" + hashlib.md5((nome + dig).encode()).hexdigest()[:10],
                "nome": nome, "cat": cat, "tel": tel, "end": end,
                "nota": nota, "aval": int(aval) if aval.isdigit() else None,
                "ig": ig, "fb": fb, "site": site, "email": "",
                "nicho": nicho, "_porque": porque,
                "site_fraco": any(f in sl for f in SITE_FRACO),
                "wa_no_site": False, "crm": [], "site_respondeu": None,
            })
    return out


def pontuar(l):
    """Prioridade 1..4 e o gancho da ligação."""
    aval = l["aval"] or 0
    celular = bool(re.search(r"\(\d\d\)\s?9", l["tel"]))
    sem_site = not l["site"] or l["site_fraco"]
    # Site listado que não abriu no robô: pode ser só bloqueio. Não é "sem site".
    site_mudo = bool(l["site"]) and not l["site_fraco"] and l["site_respondeu"] is False
    tem_crm = bool(l["crm"])

    if sem_site and aval >= 25:
        p = 1
        onde = "só link na bio" if l["site_fraco"] else "só Instagram" if l["ig"] else "nenhum site"
        g = (f"{aval} avaliações e {onde}: "
             f"a demanda existe, falta onde converter. Site + WhatsApp oficial com CRM.")
    elif site_mudo:
        p = 2 if aval >= 80 else 3
        g = "Tem site listado que não abriu na checagem: conferir antes. Se estiver fora do ar, é o gancho; senão pitch CRM + API oficial."
    elif sem_site:
        p = 2
        g = "Sem site próprio. Oferecer site + WhatsApp oficial para começar a captar."
    elif not tem_crm and aval >= 80:
        p = 2
        g = (f"Site ok, {aval} avaliações e nenhum CRM visível: o gargalo é o WhatsApp. "
             f"Pitch: API oficial + CRM, sem risco de número banido.")
    elif not tem_crm:
        p = 3
        g = "Tem site, sem automação visível. Pitch CRM + API oficial do WhatsApp."
    else:
        p = 4
        g = f"Já usa {', '.join(l['crm'])}. Só ligar se for troca: comparar custo e API oficial."
    if not celular and p < 4:
        g += " Fixo no Maps: pedir o WhatsApp de quem decide."
    if l["_porque"]:
        g += f" ({l['_porque']})"
    l["prioridade"] = p
    l["gancho"] = g
    return l


def main():
    rapido = "--rapido" in sys.argv
    leads = ler_brutos()
    print(f"{len(leads)} leads únicos (sem redes, sem fechados)")
    if not rapido:
        com_site = [l for l in leads if l["site"]]
        print(f"visitando {len(com_site)} sites…")
        with ThreadPoolExecutor(12) as ex:
            list(ex.map(enriquecer, com_site))
    leads = [pontuar(l) for l in leads]
    leads.sort(key=lambda l: (l["prioridade"], -(l["aval"] or 0)))

    saida = [{k: v for k, v in l.items() if not k.startswith("_")} for l in leads]
    (PASTA / "leads.json").write_text(json.dumps(saida, ensure_ascii=False, indent=1), encoding="utf-8")

    cols = ["prioridade", "nicho", "nome", "tel", "ig", "site", "email", "nota", "aval", "cat", "end", "gancho"]
    cab = ["Prioridade", "Nicho", "Nome", "Telefone", "Instagram", "Site", "E-mail", "Nota", "Avaliações",
           "Categoria", "Endereço", "Gancho da ligação"]
    with open(PASTA / "coldcall_uberlandia.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(cab)
        for l in leads:
            w.writerow([l[c] if l[c] is not None else "" for c in cols])
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Cold call UDI"
        ws.append(cab)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F2937")
        cores = {1: "D1FAE5", 2: "FEF3C7", 3: "F3F4F6", 4: "FEE2E2"}
        for l in leads:
            ws.append([l[c] if l[c] is not None else "" for c in cols])
            ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=cores[l["prioridade"]])
            for col in (5, 6):
                v = ws.cell(ws.max_row, col).value
                if v:
                    ws.cell(ws.max_row, col).hyperlink = v
        for col, larg in zip("ABCDEFGHIJKL", [10, 18, 38, 17, 34, 30, 30, 6, 11, 22, 36, 90]):
            ws.column_dimensions[col].width = larg
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(PASTA / "coldcall_uberlandia.xlsx")
    except ImportError:
        print("openpyxl ausente: só CSV")

    from collections import Counter
    print("por prioridade:", dict(sorted(Counter(l["prioridade"] for l in leads).items())))
    print("por nicho:", dict(Counter(l["nicho"] for l in leads).most_common()))
    print("com instagram:", sum(1 for l in leads if l["ig"]), "| com e-mail:", sum(1 for l in leads if l["email"]),
          "| com CRM detectado:", sum(1 for l in leads if l["crm"]))


if __name__ == "__main__":
    main()
